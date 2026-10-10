"""Makes the three small messy workbooks that the data prep tests read.

    python3 make_prep_fixtures.py

Needs openpyxl. The .xlsx files are committed, so the tests do not need Python.
The numbers are worked out by hand in ../prep.rs.
"""

import datetime as dt
import os

from openpyxl import Workbook

here = os.path.dirname(os.path.abspath(__file__))


def day(d):
    return dt.datetime(2026, 1, d)


# 1. A title block above the header and a footer under the data.
wb = Workbook()
ws = wb.active
ws.title = "Sales"
ws["A1"] = "Laporan Penjualan Toko Senja"
ws["A2"] = "Periode: Januari 2026"
ws.append([])
ws.append(["Date", "Region", "Product", "Amount", "Qty"])
rows = [
    (day(2), "North", "Arabica", 100, 2),
    (day(3), "South", "Robusta", 50, 1),
    (day(5), "North", "Arabica", 200, 4),
    (day(8), "South", "Liberica", 30, 1),
    (day(9), "North", "Robusta", 70, 2),
]
for r in rows:
    ws.append(list(r))
ws.append([])
ws.append(["Total", None, None, 450, 10])
ws.append(["Sumber: data internal"])
wb.save(os.path.join(here, "title_footer.xlsx"))

# 2. Merged year headers over month headers, and a region that is only written on its first row.
wb = Workbook()
ws = wb.active
ws.title = "Revenue"
ws["C1"] = 2025
ws.merge_cells("C1:D1")
ws["E1"] = 2026
ws.merge_cells("E1:F1")
ws.append(["Region", "Channel", "Jan", "Feb", "Jan", "Feb"])
ws.append(["North", "Online", 10, 20, 30, 40])
ws.append([None, "Store", 11, 21, 31, 41])
ws.append([None, "Partner", 12, 22, 32, 42])
ws.append(["South", "Online", 5, 6, 7, 8])
ws.append([None, "Store", 15, 16, 17, 18])
ws.append([None, "Partner", 25, 26, 27, 28])
wb.save(os.path.join(here, "merged_headers.xlsx"))

# 3. Months across the top: one number column for each month.
wb = Workbook()
ws = wb.active
ws.title = "Budget"
ws.append(["Department", "Owner", "Jan-26", "Feb-26", "Mar-26", "Apr-26"])
ws.append(["Kitchen", "Sari", 100, 110, 120, 130])
ws.append(["Bar", "Budi", 50, 60, None, 80])
ws.append(["Office", "Dewi", 10, 10, 10, 10])
wb.save(os.path.join(here, "months_across.xlsx"))
